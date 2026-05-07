"""Genera docs/data-dictionary.xlsx con el diccionario de datos de Quiniela 2026."""
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── Paleta Neon Tokyo ─────────────────────────────────────────────────────────
C_BG_DARK   = "0E0D14"   # fondo general
C_BG_CARD   = "13111F"   # fondo card
C_BG_HEADER = "1A1830"   # fondo fila encabezado de tabla
C_ACCENT    = "FF2D78"   # rosa neón (PK / título)
C_TEAL      = "00F5C4"   # teal neón (FK)
C_PURPLE    = "7B2FFB"   # púrpura (UNIQUE)
C_TEXT      = "E8E0F4"   # texto principal
C_MUTED     = "9B8EC4"   # texto secundario
C_WHITE     = "FFFFFF"
C_STRIP     = "16142B"   # fila alternada

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(hex_color=C_TEXT, bold=False, size=10):
    return Font(color=hex_color, bold=bold, size=size, name="Segoe UI")

thin = Side(style="thin", color="2A2550")
thick = Side(style="medium", color=C_ACCENT)
border_row  = Border(left=thin, right=thin, top=thin, bottom=thin)
border_head = Border(left=thick, right=thick, top=thick, bottom=thick)

COLS = ["Columna", "Tipo SQL", "Nullable", "Restricciones", "Descripción"]
COL_W = [26, 14, 10, 28, 48]

# ── Datos ─────────────────────────────────────────────────────────────────────
TABLES = [
    {
        "name": "users",
        "desc": "Jugadores y administradores. El flujo de acceso es: Pending → Active (o Rejected) mediante aprobación admin.",
        "rows": [
            ("id",             "INTEGER", "NO",  "PK · autoincrement",                     "Identificador único del usuario."),
            ("username",       "TEXT",    "NO",  "UNIQUE · NOT NULL",                      "Nombre de usuario para login. Inmutable una vez creado."),
            ("password_hash",  "TEXT",    "NO",  "NOT NULL",                               "Hash bcrypt de la contraseña."),
            ("is_admin",       "INTEGER", "NO",  "NOT NULL · default 0",                   "0 = jugador normal · 1 = administrador."),
            ("status",         "TEXT",    "NO",  "NOT NULL · default 'Pending'",           "Estado de la cuenta: Pending | Active | Rejected."),
            ("total_points",   "INTEGER", "NO",  "NOT NULL · default 0",                   "Suma acumulada de puntos obtenidos en predicciones y bonus."),
            ("exact_scores",   "INTEGER", "NO",  "NOT NULL · default 0",                   "Contador de marcadores exactos acertados. Usado como desempate en el ranking."),
            ("created_at",     "DATETIME","NO",  "NOT NULL · default utcnow",              "Fecha y hora UTC de creación del registro."),
        ],
    },
    {
        "name": "matches",
        "desc": "Partidos del torneo. Los partidos de Fase de Grupos se crean vía seed_matches.py; los eliminatorios los crea el Admin. El estado lógico (Open/Locked/Finished/Pending Teams) se deriva en runtime por time_utils.py y nunca se persiste.",
        "rows": [
            ("id",                    "INTEGER", "NO",  "PK · autoincrement",            "Identificador único del partido."),
            ("home_team",             "TEXT",    "SÍ",  "nullable",                       "Nombre del equipo local. NULL en eliminatorias hasta que el Admin asigne equipos."),
            ("away_team",             "TEXT",    "SÍ",  "nullable",                       "Nombre del equipo visitante. NULL en eliminatorias hasta asignación."),
            ("home_team_code",        "TEXT",    "SÍ",  "nullable",                       "Código ISO 3166-1 alpha-2 del equipo local (ej. 'mx'). Usado para cargar la bandera desde flagcdn.com."),
            ("away_team_code",        "TEXT",    "SÍ",  "nullable",                       "Código ISO 3166-1 alpha-2 del equipo visitante."),
            ("home_team_placeholder", "TEXT",    "SÍ",  "nullable",                       "Texto descriptivo del slot local en eliminatorias (ej. '1A vs 2B')."),
            ("away_team_placeholder", "TEXT",    "SÍ",  "nullable",                       "Texto descriptivo del slot visitante en eliminatorias."),
            ("start_time",            "DATETIME","NO",  "NOT NULL",                       "Fecha y hora UTC de inicio del partido. Referencia para calcular el estado lógico."),
            ("phase",                 "TEXT",    "NO",  "NOT NULL",                       "Fase del torneo: Groups | R32 | R16 | QF | SF | ThirdPlace | Final."),
            ("group_name",            "TEXT",    "SÍ",  "nullable",                       "Grupo al que pertenece (A–L). Solo aplica a phase = Groups."),
            ("matchday",              "INTEGER", "SÍ",  "nullable",                       "Jornada dentro del grupo (1, 2 o 3). Solo aplica a phase = Groups."),
            ("venue",                 "TEXT",    "SÍ",  "nullable",                       "Sede del partido (ciudad/estadio)."),
            ("home_score_final",      "INTEGER", "SÍ",  "nullable",                       "Goles del equipo local al final del partido. NULL hasta que el Admin cargue el resultado."),
            ("away_score_final",      "INTEGER", "SÍ",  "nullable",                       "Goles del equipo visitante al final del partido. NULL hasta resultado."),
            ("status",                "TEXT",    "NO",  "NOT NULL · default 'Scheduled'", "Estado persistido del partido: Scheduled | Finished. El resto (Open, Locked, Pending Teams) es runtime."),
            ("created_at",            "DATETIME","NO",  "NOT NULL · default utcnow",      "Fecha y hora UTC de inserción del partido."),
        ],
    },
    {
        "name": "predictions",
        "desc": "Predicciones de marcador de cada usuario por partido. La lógica de upsert garantiza una única predicción por par (user_id, match_id). Solo se puede crear/editar mientras el partido esté en estado Open (más de 15 min antes del inicio).",
        "rows": [
            ("id",               "INTEGER", "NO",  "PK · autoincrement",               "Identificador único de la predicción."),
            ("user_id",          "INTEGER", "NO",  "FK → users.id · CASCADE · NOT NULL","Usuario que realizó la predicción."),
            ("match_id",         "INTEGER", "NO",  "FK → matches.id · CASCADE · NOT NULL","Partido al que corresponde la predicción."),
            ("home_score_guess", "INTEGER", "NO",  "NOT NULL",                          "Goles pronosticados para el equipo local."),
            ("away_score_guess", "INTEGER", "NO",  "NOT NULL",                          "Goles pronosticados para el equipo visitante."),
            ("points_earned",    "INTEGER", "SÍ",  "nullable",                          "Puntos obtenidos: 0, 3 (tendencia) o 5 (exacto). NULL hasta que el Admin cargue el resultado."),
            ("created_at",       "DATETIME","NO",  "NOT NULL · default utcnow",         "Fecha y hora UTC de creación."),
            ("updated_at",       "DATETIME","NO",  "NOT NULL · default utcnow",         "Fecha y hora UTC de la última edición (upsert)."),
            ("—",                "—",       "—",   "UNIQUE (user_id, match_id)",        "Restricción compuesta: un usuario no puede tener dos predicciones para el mismo partido."),
        ],
    },
    {
        "name": "bonus_predictions",
        "desc": "Predicciones bonus (Campeón del Mundo y Bota de Oro). Se cierran automáticamente cuando el tiempo de referencia supera el start_time del primer partido en la DB. Los puntos se asignan manualmente por el Admin.",
        "rows": [
            ("id",               "INTEGER", "NO",  "PK · autoincrement",               "Identificador único."),
            ("user_id",          "INTEGER", "NO",  "FK → users.id · CASCADE · NOT NULL","Usuario que realizó la predicción bonus."),
            ("question_type",    "TEXT",    "NO",  "NOT NULL",                          "Tipo de pregunta bonus: Champion | TopScorer."),
            ("prediction_text",  "TEXT",    "NO",  "NOT NULL",                          "Texto libre con la predicción (nombre del equipo o jugador)."),
            ("points_earned",    "INTEGER", "SÍ",  "nullable",                          "Puntos asignados por el Admin: 20 (Champion) o 15 (TopScorer). NULL hasta validación."),
            ("created_at",       "DATETIME","NO",  "NOT NULL · default utcnow",         "Fecha y hora UTC de creación."),
            ("—",                "—",       "—",   "UNIQUE (user_id, question_type)",   "Un usuario solo puede tener una predicción por tipo de bonus."),
        ],
    },
]

# ── Reglas de puntuación (hoja extra) ─────────────────────────────────────────
SCORING = [
    ("Tendencia correcta",  "+3 pts", "El usuario predijo quién gana o que hay empate (1X2) y coincide con el resultado final."),
    ("Marcador exacto",     "+5 pts", "El marcador local-visitante predicho es idéntico al resultado final. Incluye el bono de tendencia (+2 adicionales sobre los +3 base)."),
    ("Tendencia incorrecta", "0 pts", "La predicción no coincide con el resultado (ni en ganador ni en empate)."),
    ("Bonus Campeón",       "+20 pts","Predicción de texto libre. Puntos asignados manualmente por el Admin cuando finaliza el torneo."),
    ("Bonus Bota de Oro",   "+15 pts","Predicción de texto libre. Puntos asignados manualmente por el Admin cuando finaliza el torneo."),
    ("Desempate ranking",   "—",      "En caso de igualdad de total_points: 1º exact_scores DESC · 2º username ASC."),
]

# ── Build workbook ─────────────────────────────────────────────────────────────

wb = Workbook()
wb.remove(wb.active)  # quita la hoja default

# ── Índice ────────────────────────────────────────────────────────────────────
ws_idx = wb.create_sheet("Índice")
ws_idx.sheet_view.showGridLines = False
ws_idx.column_dimensions["A"].width = 6
ws_idx.column_dimensions["B"].width = 28
ws_idx.column_dimensions["C"].width = 62

# título
ws_idx.merge_cells("A1:C1")
c = ws_idx["A1"]
c.value = "Quiniela 2026 — Diccionario de Datos"
c.font = Font(color=C_ACCENT, bold=True, size=16, name="Segoe UI")
c.fill = fill(C_BG_DARK)
c.alignment = Alignment(horizontal="center", vertical="center")
ws_idx.row_dimensions[1].height = 36

ws_idx.merge_cells("A2:C2")
c = ws_idx["A2"]
c.value = "Base de datos SQLite · ORM SQLAlchemy · Backend FastAPI"
c.font = Font(color=C_MUTED, size=10, name="Segoe UI")
c.fill = fill(C_BG_DARK)
c.alignment = Alignment(horizontal="center")
ws_idx.row_dimensions[2].height = 20

ws_idx.row_dimensions[3].height = 8

# encabezados tabla índice
for col, (txt, w) in enumerate(zip(["#", "Tabla", "Descripción"], [6, 28, 62]), 1):
    c = ws_idx.cell(row=4, column=col, value=txt)
    c.font = Font(color=C_WHITE, bold=True, size=10, name="Segoe UI")
    c.fill = fill(C_BG_HEADER)
    c.alignment = Alignment(horizontal="center" if col == 1 else "left", vertical="center")
    c.border = border_row

ws_idx.row_dimensions[4].height = 20

for i, t in enumerate(TABLES, 1):
    r = i + 4
    bg = C_BG_CARD if i % 2 == 0 else C_STRIP
    for col, val in enumerate([i, t["name"], t["desc"]], 1):
        c = ws_idx.cell(row=r, column=col, value=val)
        c.font = Font(color=C_ACCENT if col == 2 else C_TEXT, bold=(col==2), size=10, name="Segoe UI")
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center" if col==1 else "left", vertical="top", wrap_text=(col==3))
        c.border = border_row
    ws_idx.row_dimensions[r].height = 42

# ── Una hoja por tabla ────────────────────────────────────────────────────────
for t in TABLES:
    ws = wb.create_sheet(t["name"])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    for ci, w in enumerate(COL_W, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    # título hoja
    ws.merge_cells(f"A1:{get_column_letter(len(COLS))}1")
    c = ws["A1"]
    c.value = f"Tabla: {t['name']}"
    c.font = Font(color=C_ACCENT, bold=True, size=14, name="Segoe UI")
    c.fill = fill(C_BG_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    # descripción
    ws.merge_cells(f"A2:{get_column_letter(len(COLS))}2")
    c = ws["A2"]
    c.value = t["desc"]
    c.font = Font(color=C_MUTED, size=9, name="Segoe UI")
    c.fill = fill(C_BG_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 40

    # encabezados columnas
    for ci, col_name in enumerate(COLS, 1):
        c = ws.cell(row=3, column=ci, value=col_name)
        c.font = Font(color=C_WHITE, bold=True, size=10, name="Segoe UI")
        c.fill = fill(C_BG_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_head
    ws.row_dimensions[3].height = 22

    # filas
    for ri, row in enumerate(t["rows"], 4):
        col_name, tipo, nullable, restricciones, desc = row
        bg = C_BG_CARD if ri % 2 == 0 else C_STRIP

        is_pk   = "PK" in restricciones
        is_fk   = "FK" in restricciones
        is_uniq = restricciones.startswith("UNIQUE") or (col_name == "—")

        name_color = C_ACCENT if is_pk else (C_TEAL if is_fk else (C_PURPLE if is_uniq else C_TEXT))

        vals = [col_name, tipo, nullable, restricciones, desc]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill = fill(bg)
            c.border = border_row
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if ci == 1:
                c.font = Font(color=name_color, bold=True, size=10, name="Segoe UI")
            elif ci == 3:
                nullable_color = "FF6B6B" if val == "SÍ" else "6BFF9E"
                c.font = Font(color=nullable_color, size=10, name="Segoe UI")
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                c.font = Font(color=C_TEXT if ci != 4 else C_MUTED, size=10, name="Segoe UI")
        ws.row_dimensions[ri].height = 36

# ── Hoja de puntuación ────────────────────────────────────────────────────────
ws_sc = wb.create_sheet("Puntuación")
ws_sc.sheet_view.showGridLines = False

sc_cols = ["Concepto", "Puntos", "Descripción"]
sc_widths = [26, 12, 60]
for ci, w in enumerate(sc_widths, 1):
    ws_sc.column_dimensions[get_column_letter(ci)].width = w

ws_sc.merge_cells("A1:C1")
c = ws_sc["A1"]
c.value = "Reglas de Puntuación"
c.font = Font(color=C_ACCENT, bold=True, size=14, name="Segoe UI")
c.fill = fill(C_BG_DARK)
c.alignment = Alignment(horizontal="left", vertical="center")
ws_sc.row_dimensions[1].height = 32

for ci, col_name in enumerate(sc_cols, 1):
    c = ws_sc.cell(row=2, column=ci, value=col_name)
    c.font = Font(color=C_WHITE, bold=True, size=10, name="Segoe UI")
    c.fill = fill(C_BG_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_head
ws_sc.row_dimensions[2].height = 22

for ri, (concepto, pts, desc) in enumerate(SCORING, 3):
    bg = C_BG_CARD if ri % 2 == 0 else C_STRIP
    pt_color = C_TEAL if pts.startswith("+") else (C_MUTED if pts == "—" else "FF6B6B")
    for ci, (val, fc) in enumerate(zip([concepto, pts, desc], [C_TEXT, pt_color, C_MUTED]), 1):
        c = ws_sc.cell(row=ri, column=ci, value=val)
        c.font = Font(color=fc, bold=(ci==2), size=10, name="Segoe UI")
        c.fill = fill(bg)
        c.alignment = Alignment(horizontal="center" if ci==2 else "left", vertical="top", wrap_text=True)
        c.border = border_row
    ws_sc.row_dimensions[ri].height = 32

# ── Guardar ───────────────────────────────────────────────────────────────────
OUT = "docs/data-dictionary.xlsx"
wb.save(OUT)
print(f"✓ {OUT} generado correctamente.")
